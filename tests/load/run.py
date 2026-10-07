"""Real HTTP upload/poll/download benchmark; synthetic blank forms only.

Run against tests/load/compose.yml. Only dummy credentials for the local API simulator are supplied.
"""
import argparse
import asyncio
import io
import json
import math
import time
from pathlib import Path

import httpx
from docx import Document
from reportlab.pdfgen import canvas


TEXT = "Please write your name in the box. Give the court your mailing address. Sign this form and keep a copy."


def pdf(pages=1, fields=True):
    stream = io.BytesIO()
    doc = canvas.Canvas(stream, pagesize=(612, 792))
    for page in range(pages):
        doc.setTitle("Synthetic court form")
        doc.drawString(50, 745, "Request for a court date")
        for line in range(12):
            doc.drawString(50, 710 - 20 * line, TEXT[:95])
        if fields:
            doc.acroForm.textfield(name=f"name_{page}", tooltip="Your name", x=50, y=380, width=200, height=20)
        doc.showPage()
    doc.save()
    return stream.getvalue()


def fixtures():
    word = Document()
    word.add_heading("Request for a court date", 0)
    for _ in range(12):
        word.add_paragraph(TEXT)
    stream = io.BytesIO()
    word.save(stream)
    small = pdf()
    return {
        "small": ("small.pdf", small),
        "unfielded": ("unfielded.pdf", pdf(fields=False)),
        "medium": ("medium.pdf", pdf(10)),
        "large": ("large.pdf", pdf(50)),
        "word": ("form.docx", stream.getvalue()),
        "boundary": ("boundary.pdf", small + b"\n%" + b" " * (20 * 1024 * 1024 - len(small) - 2)),
        "oversize": ("incident.pdf", small + b" " * (62 * 1024 * 1024)),
    }


def distribution(values):
    values = sorted(values)
    if not values:
        return {}
    return {"count": len(values), "p50": round(values[math.ceil(len(values) * .5)-1], 3),
            "p95": round(values[math.ceil(len(values) * .95)-1], 3), "max": round(values[-1], 3)}


async def run(args):
    data = fixtures()
    data["oversize_stream"] = data["oversize"]
    api_before = None
    try:
        api_before = httpx.get(args.simulator + "/metrics").json()
    except httpx.HTTPError:
        pass
    health = []
    done = asyncio.Event()
    results = []
    start = time.monotonic()
    async with httpx.AsyncClient(base_url=args.url, timeout=20, limits=httpx.Limits(max_connections=50)) as client:
        async def probe():
            while not done.is_set():
                begin = time.monotonic()
                try:
                    response = await client.get("/healthz", timeout=5)
                    status = response.status_code
                except Exception as exc:
                    status = type(exc).__name__
                health.append({"elapsed": round(time.monotonic()-start, 3), "seconds": time.monotonic()-begin, "status": status})
                await asyncio.sleep(.5)

        async def upload(case):
            filename, content = data[case]
            begin = time.monotonic()
            record = {"case": case, "bytes": len(content), "stages": []}
            try:
                if case == "oversize_stream":
                    async def chunks():
                        yield b'--benchmark\r\nContent-Disposition: form-data; name="terms_version"\r\n\r\n2026-10-07\r\n'
                        yield b'--benchmark\r\nContent-Disposition: form-data; name="file"; filename="stream.pdf"\r\nContent-Type: application/pdf\r\n\r\n'
                        for offset in range(0, len(content), 64 * 1024):
                            yield content[offset:offset + 64 * 1024]
                        yield b'\r\n--benchmark--\r\n'
                    response = await client.post("/", content=chunks(), headers={"Content-Type": "multipart/form-data; boundary=benchmark"})
                else:
                    response = await client.post("/", data={"terms_version": "2026-10-07"}, files={"file": (filename, content)})
                record.update(upload_seconds=time.monotonic()-begin, upload_status=response.status_code)
                if response.status_code != 303:
                    record["outcome"] = "rejected" if response.status_code == 413 else "upload_error"
                    record["detail"] = response.text[:300]
                    return record
                location = response.headers["location"]
                record["job_id"] = location.rsplit("/", 1)[-1]
                page = await client.get(location)
                record["view_status"] = page.status_code
                while time.monotonic()-begin < args.deadline:
                    response = await client.get("/job-status/" + record["job_id"])
                    if response.status_code != 200:
                        record.update(outcome="poll_error", poll_status=response.status_code, detail=response.text[:300])
                        break
                    status = response.json()
                    stage = status.get("status_message", status["status"])
                    if not record["stages"] or stage != record["stages"][-1]["stage"]:
                        record["stages"].append({"stage": stage, "seconds": round(time.monotonic()-begin, 3)})
                    if status["status"] == "finished":
                        download = await client.get("/download/" + record["job_id"])
                        record.update(outcome="finished" if download.status_code == 200 and download.content.startswith(b"%PDF") and status["rendered_stats"] and page.status_code == 200 else "validation_error", html_bytes=len(status["rendered_stats"]), download_status=download.status_code,
                                      valid_pdf=download.content.startswith(b"%PDF"))
                        break
                    if status["status"] != "pending":
                        record.update(outcome=status["status"])
                        break
                    await asyncio.sleep(args.poll_interval)
                else:
                    record["outcome"] = "deadline"
            except Exception as exc:
                record.update(outcome="client_error", detail=f"{type(exc).__name__}: {exc}")
            finally:
                record["total_seconds"] = round(time.monotonic()-begin, 3)
                print(json.dumps(record), flush=True)
            return record

        probe_task = asyncio.create_task(probe())
        semaphore = asyncio.Semaphore(args.concurrency)
        async def limited(case):
            async with semaphore:
                return await upload(case)
        try:
            results = await asyncio.gather(*(limited(case) for case in args.cases.split(",")))
        finally:
            done.set()
            await probe_task
    report = {"profile": args.profile, "url": args.url, "concurrency": args.concurrency, "wall_seconds": round(time.monotonic()-start, 3),
              "results": results, "health_latency_seconds": distribution([h["seconds"] for h in health]),
              "health_errors": [h for h in health if h["status"] != 200],
              "successful_total_seconds": distribution([r["total_seconds"] for r in results if r["outcome"] == "finished"]),
              "health_samples": health}
    if api_before is not None:
        api_after = httpx.get(args.simulator + "/metrics").json()
        report["simulated_api"] = {"delay_seconds": api_after["delay_seconds"], "calls": api_after["calls"] - api_before["calls"]}
    Path(args.output).write_text(json.dumps(report, indent=2))
    print(json.dumps({k: v for k, v in report.items() if k not in {"health_samples", "results"}}, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--profile", default="container-limits")
    parser.add_argument("--url", default="http://127.0.0.1:18080")
    parser.add_argument("--simulator", default="http://127.0.0.1:18081")
    parser.add_argument("--cases", default="small,unfielded,word,medium,large,boundary,oversize")
    parser.add_argument("--concurrency", type=int, default=1)
    parser.add_argument("--deadline", type=int, default=600)
    parser.add_argument("--poll-interval", type=float, default=2)
    parser.add_argument("--output", required=True)
    asyncio.run(run(parser.parse_args()))
